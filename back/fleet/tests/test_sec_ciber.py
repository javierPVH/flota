"""Auditoría de ciberseguridad (2026-09): INP-1, AUTH-7, AUTH-10 e INP-7 (fleet)."""

import io
import tempfile
from datetime import date
from pathlib import Path

from django.core.exceptions import ValidationError
from django.test import TestCase, override_settings
from django.urls import reverse
from openpyxl import load_workbook
from rest_framework import status
from rest_framework.test import APITestCase

from accounts.models import Role, User, UserRole
from fleet.models import Assignment, Document, Vehicle
from fleet.models.enums import AssignmentStatus, DocumentStatus
from fleet.models.vehicle import PLATE_VALIDATOR
from fleet.services import archiver, reports

from .helpers import make_user


class Inp1XlsxFormulaTests(TestCase):
    """INP-1: el Excel no interpreta como fórmula lo que escribió un usuario."""

    FORMULA = '=HYPERLINK("http://x","y")'

    def _sheet(self, rows):
        payload = reports.to_xlsx([("Hoja", ["Notas", "Km", "Fecha"], rows)])
        return load_workbook(io.BytesIO(payload))["Hoja"]

    def test_inp1_xlsx_no_interpreta_formulas(self):
        ws = self._sheet([[self.FORMULA, 123, "2026-09-22"]])
        celda = ws["A2"]
        # Tipo cadena, no fórmula (`f`), y el texto se conserva tal cual: sin
        # el apóstrofo que añade el CSV.
        self.assertEqual(celda.data_type, "s")
        self.assertEqual(celda.value, self.FORMULA)
        # Lo que no es texto sigue viajando tipado.
        self.assertEqual(ws["B2"].value, 123)
        self.assertEqual(ws["B2"].data_type, "n")
        self.assertEqual(ws["C2"].value, "2026-09-22")

    def test_inp1_otros_prefijos_peligrosos_y_cabeceras(self):
        peligrosas = ["+cmd", "-2+3", "@SUM(A1)", "\tX", "=1+1"]
        ws = self._sheet([[valor, 1, ""] for valor in peligrosas])
        for fila, valor in enumerate(peligrosas, start=2):
            celda = ws.cell(row=fila, column=1)
            self.assertEqual(celda.data_type, "s", valor)
            self.assertEqual(celda.value, valor)
        # La cabecera sigue en negrita: forzar el tipo no se lleva el estilo.
        self.assertTrue(ws["A1"].font.bold)
        self.assertEqual(ws["A1"].value, "Notas")

    def test_inp1_el_informe_real_lleva_el_texto_como_cadena(self):
        admin = make_user("admin", Role.ADMIN)
        vehicle = Vehicle.objects.create(plate="INP1AAA", brand="a", model="b")
        Document.objects.create(vehicle=vehicle, type="other", notes=self.FORMULA)
        _, _, payload = reports.render("documents", admin, "xlsx")
        ws = load_workbook(io.BytesIO(payload)).worksheets[0]
        notas = [c for c in ws[2] if c.value == self.FORMULA]
        self.assertEqual(len(notas), 1)
        self.assertEqual(notas[0].data_type, "s")


class Auth7RestoreSuperuserTests(APITestCase):
    """AUTH-7: un admin normal no reactiva a un superusuario desde erratas."""

    def setUp(self):
        self.admin = make_user("admin", Role.ADMIN)
        self.root = User.objects.create_superuser(
            username="root", email="root@example.com", password="test-pass-123"
        )
        UserRole.objects.create(user=self.root, role=Role.ADMIN)
        self.root.is_active = False
        self.root.save(update_fields=["is_active"])
        self.otro_root = User.objects.create_superuser(
            username="root2", email="root2@example.com", password="test-pass-123"
        )
        UserRole.objects.create(user=self.otro_root, role=Role.ADMIN)
        self.url = reverse("erratas-restore")

    def test_auth7_admin_normal_recibe_403_y_no_reactiva(self):
        self.client.force_authenticate(self.admin)
        resp = self.client.post(self.url, {"type": "users", "id": self.root.pk}, format="json")
        self.assertEqual(resp.status_code, status.HTTP_403_FORBIDDEN)
        self.root.refresh_from_db()
        self.assertFalse(self.root.is_active)

    def test_auth7_superusuario_si_puede_restaurar(self):
        self.client.force_authenticate(self.otro_root)
        resp = self.client.post(self.url, {"type": "users", "id": self.root.pk}, format="json")
        self.assertEqual(resp.status_code, status.HTTP_200_OK)
        self.root.refresh_from_db()
        self.assertTrue(self.root.is_active)

    def test_auth7_admin_normal_sigue_restaurando_usuarios_corrientes(self):
        conductor = make_user("carlos", Role.DRIVER)
        conductor.is_active = False
        conductor.save(update_fields=["is_active"])
        self.client.force_authenticate(self.admin)
        resp = self.client.post(self.url, {"type": "users", "id": conductor.pk}, format="json")
        self.assertEqual(resp.status_code, status.HTTP_200_OK)
        conductor.refresh_from_db()
        self.assertTrue(conductor.is_active)


class Auth10InactiveMediaTests(APITestCase):
    """AUTH-10: el binario de un documento desactivado deja de servirse por /media."""

    PATH = "documents/2026/09/foto-inactiva.jpg"

    def setUp(self):
        self.admin = make_user("media-admin", Role.ADMIN)
        self.driver = make_user("media-driver", Role.DRIVER)
        self.supervisor = make_user("media-super", Role.SUPERVISOR)
        self.vehicle = Vehicle.objects.create(
            plate="AUTH10A", brand="a", model="b", supervisor=self.supervisor
        )
        Assignment.objects.create(
            vehicle=self.vehicle,
            driver=self.driver,
            start_date=date(2026, 1, 1),
            status=AssignmentStatus.ACCEPTED,
        )
        # Compartido y con el conductor de responsable: todo el mundo del
        # ámbito lo lee mientras esté vivo.
        self.document = Document.objects.create(
            vehicle=self.vehicle,
            type="damage_photos",
            file=self.PATH,
            responsible=self.driver,
            shared_read=True,
        )

    @override_settings(DEBUG=False)
    def test_auth10_documento_vivo_se_sirve(self):
        self.client.force_authenticate(self.driver)
        self.assertEqual(self.client.get(f"/media/{self.PATH}").status_code, status.HTTP_200_OK)

    @override_settings(DEBUG=False)
    def test_auth10_documento_desactivado_es_404_para_el_campo(self):
        self.document.deactivate(by=self.admin, reason="errata")
        for user in (self.driver, self.supervisor):
            self.client.force_authenticate(user)
            resp = self.client.get(f"/media/{self.PATH}")
            self.assertEqual(resp.status_code, status.HTTP_404_NOT_FOUND, user.username)

    @override_settings(DEBUG=False)
    def test_auth10_la_gestion_sigue_viendolo_desde_erratas(self):
        self.document.deactivate(by=self.admin, reason="errata")
        self.client.force_authenticate(self.admin)
        self.assertEqual(self.client.get(f"/media/{self.PATH}").status_code, status.HTTP_200_OK)


class Inp7PlateValidatorTests(APITestCase):
    """INP-7: la matrícula tiene forma de matrícula, no de ruta."""

    def test_inp7_matriculas_validas(self):
        for plate in ("1234ABC", "M-1234-AB", "1234 ABC", "P-2", "GAP-0001", "0000ZZZ"):
            PLATE_VALIDATOR(plate)  # no lanza

    def test_inp7_matriculas_invalidas(self):
        for plate in (
            "../../x",
            "..",
            "1234abc",
            " 1234ABC",
            "1234ABC ",
            "12",
            "A/B",
            "1234.ABC",
            "A" * 16,
        ):
            with self.assertRaises(ValidationError, msg=plate):
                PLATE_VALIDATOR(plate)

    def test_inp7_full_clean_rechaza_la_ruta(self):
        vehicle = Vehicle(plate="../../fuera", brand="a", model="b")
        with self.assertRaises(ValidationError) as ctx:
            vehicle.full_clean()
        self.assertIn("plate", ctx.exception.message_dict)

    def test_inp7_la_api_devuelve_400(self):
        admin = make_user("admin", Role.ADMIN)
        self.client.force_authenticate(admin)
        resp = self.client.post(
            reverse("vehicle-list"), {"plate": "../../pwn", "brand": "a", "model": "b"}
        )
        self.assertEqual(resp.status_code, status.HTTP_400_BAD_REQUEST)
        # El handler de errores del proyecto envuelve los de campo en `errors`.
        self.assertIn("plate", resp.data["errors"])
        self.assertFalse(Vehicle.objects.filter(brand="a").exists())


class Inp7LocalArchiverEscapeTests(TestCase):
    """INP-7: el archivador `local` no escribe fuera de su carpeta base."""

    def setUp(self):
        # Por el ORM el validador no corre: es justo el caso que el archivador
        # tiene que parar por su cuenta.
        self.vehicle = Vehicle.objects.create(plate="../../fuera", brand="a", model="b")

    def test_inp7_archive_lanza_valueerror_y_no_escribe_nada(self):
        doc = Document.objects.create(vehicle=self.vehicle, type="insurance")
        with tempfile.TemporaryDirectory() as tmp:
            base = Path(tmp) / "archivo"
            base.mkdir()
            local = archiver.LocalArchiver(base)
            with self.assertRaises(ValueError):
                local.archive(doc)
            with self.assertRaises(ValueError):
                local.ensure_folder(self.vehicle)
            # Nada ha aparecido por encima de la base.
            self.assertFalse((Path(tmp).parent / "fuera").exists())
            self.assertEqual(list(Path(tmp).iterdir()), [base])
        self.vehicle.refresh_from_db()
        self.assertEqual(self.vehicle.drive_folder_url, "")

    def test_inp7_archive_document_deja_el_documento_pendiente(self):
        doc = Document.objects.create(vehicle=self.vehicle, type="insurance")
        with tempfile.TemporaryDirectory() as tmp:
            archiver.archive_document(doc, archiver=archiver.LocalArchiver(tmp))
        self.assertEqual(doc.status, DocumentStatus.PENDING_ARCHIVE)
        self.assertEqual(doc.drive_url, "")

    def test_inp7_una_matricula_normal_sigue_archivando(self):
        normal = Vehicle.objects.create(plate="INP7OK1", brand="a", model="b")
        doc = Document.objects.create(vehicle=normal, type="insurance")
        with tempfile.TemporaryDirectory() as tmp:
            archiver.archive_document(doc, archiver=archiver.LocalArchiver(tmp))
            self.assertEqual(doc.status, DocumentStatus.VALID)
            self.assertIn("/Veh%C3%ADculos/INP7OK1/", doc.drive_url)
